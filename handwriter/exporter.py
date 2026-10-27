from __future__ import annotations

from pathlib import Path

from docx import Document
from docx.shared import Cm, Pt

from .data import DocumentProject


def validate_export(project: DocumentProject) -> None:
    if not project.pages:
        raise ValueError("任务没有页面。")
    problems = []
    for page in project.pages:
        if page.status not in {"recognized", "review"}:
            problems.append(f"第 {page.index + 1} 页尚未完成")
            continue
        if any(line.status not in {"recognized", "review"} for line in page.lines):
            problems.append(f"第 {page.index + 1} 页有未识别或已修改的行框")
    if problems:
        raise ValueError("请先完成识别后再导出：\n" + "\n".join(problems[:8]))


def export_txt(project: DocumentProject, path: str | Path) -> Path:
    validate_export(project)
    destination = Path(path)
    sections = []
    for page in sorted(project.pages, key=lambda page: page.index):
        text = "\n".join(line.text for line in page.lines)
        sections.append(f"—— 第 {page.index + 1} 页 ——\n{text}")
    temporary = destination.with_suffix(destination.suffix + ".tmp")
    temporary.write_text("\n\n".join(sections) + "\n", encoding="utf-8")
    temporary.replace(destination)
    return destination


def export_docx(project: DocumentProject, path: str | Path) -> Path:
    validate_export(project)
    destination = Path(path)
    document = Document()
    document.core_properties.title = project.title
    document.core_properties.subject = "英文手写稿转写（本地识别与人工校对）"
    section = document.sections[0]
    section.top_margin = section.bottom_margin = Cm(2)
    section.left_margin = section.right_margin = Cm(2)
    normal = document.styles["Normal"]
    normal.font.name = "Calibri"
    normal.font.size = Pt(11)
    normal.paragraph_format.space_after = Pt(0)
    for position, page in enumerate(sorted(project.pages, key=lambda page: page.index)):
        if position:
            document.add_page_break()
        heading = document.add_paragraph(f"第 {page.index + 1} 页", style="Caption")
        heading.paragraph_format.space_after = Pt(8)
        paragraph = document.add_paragraph()
        for index, line in enumerate(page.lines):
            if index:
                paragraph.add_run().add_break()
            paragraph.add_run(line.text)
    temporary = destination.with_suffix(destination.suffix + ".tmp")
    document.save(str(temporary))
    temporary.replace(destination)
    return destination
